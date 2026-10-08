from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
import re
import unicodedata
from typing import BinaryIO, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


SECONDS_PER_DAY = 86_400
NUR_SCALE = 100_000

COMMERCIAL_REGION_HEADER = "Region_commerciale"
COMMERCIAL_REGION_COLUMN_INDEX = 10  # colonne K
SITE_NAME_COLUMN_INDEX = 2  # colonne C
SITE_CODE_COLUMN_INDEX = 3  # colonne D
SITE_CODE_HEADER = "Site_Code"

REGION_ALIASES = {
    "KINSHASA": "Kinshasa",
    "GRAND KATANGA": "Grand Katanga",
    "KATANGA": "Grand Katanga",
    "GRAND KIVU": "Grand Kivu",
    "KIVU": "Grand Kivu",
    "GRAND KASAI": "Grand Kasai",
    "GRAND KASAÏ": "Grand Kasai",
    "KASAI": "Grand Kasai",
    "KASAÏ": "Grand Kasai",
    "KASAI BANDUNDU": "Grand Kasai",
    "KASAÏ BANDUNDU": "Grand Kasai",
    "KONGO CENTRAL": "Kongo Central",
    "BAS CONGO": "Kongo Central",
    "GRAND NORD": "Grand Nord",
}

REGION_NUR_THRESHOLDS = {
    "Kinshasa": 800,
    "Grand Katanga": 600,
    "Grand Kasai": 500,
    "Kongo Central": 300,
    "Grand Kivu": 1000,
    "Grand Nord": 1000,
}

REGION_ORDER = [
    "Kinshasa",
    "Grand Katanga",
    "Grand Kasai",
    "Kongo Central",
    "Grand Kivu",
    "Grand Nord",
]

NUR_SHEETS = {
    "2G": {
        "date": "Date",
        "site": "Site Name",
        "cell": "Cell Name",
        "unavailability": "R373:Cell Out-of-Service Duration(s)",
        "suffix": r"_GSM$",
    },
    "3G": {
        "date": "Date",
        "site": "NODEBNAME",
        "cell": "Cell Name",
        "unavailability": "VS.Cell.UnavailTime.Sys(s)",
        "suffix": r"_UMTS$",
    },
    "4G": {
        "date": "Date",
        "site": "eNodeB Name",
        "cell": "Cell Name",
        "unavailability": "L.Cell.Unavail.Dur.Sys(s)",
        "suffix": None,
    },
}

DATA_COLUMNS = [
    "date",
    "technology",
    "site_code",
    "cell_name",
    "unavailability_seconds",
    "source_row",
]


@dataclass
class ImportResult:
    data: list[dict]
    rejected: list[dict]
    eligible_site_count: int
    excluded_vsat_site_count: int
    total_nur_rows: int
    database_sheet_name: str


@dataclass
class CalculationResult:
    cells: list[dict]
    site_technology: list[dict]
    sites: list[dict]
    technologies: list[dict]
    regions: list[dict]
    daily: list[dict]
    network: list[dict]
    worst_sites: list[dict]
    sites_above_threshold: list[dict]
    start_date: date
    end_date: date
    period_days: int
    selected_rows: int


def _rewind(source: str | Path | BinaryIO) -> None:
    if hasattr(source, "seek"):
        source.seek(0)


def _open_workbook(source: str | Path | BinaryIO):
    _rewind(source)
    return load_workbook(source, read_only=True, data_only=True)


def _normalise_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def normalise_site_code(value: object, suffix_pattern: str | None = None) -> str:
    code = _normalise_text(value).upper()
    if suffix_pattern:
        code = re.sub(suffix_pattern, "", code, flags=re.IGNORECASE)
    return code


def _fold_region_key(value: object) -> str:
    text = unicodedata.normalize("NFKD", _normalise_text(value))
    return "".join(char for char in text if not unicodedata.combining(char)).upper()


def classify_region(value: object, fallback: str = "Non renseignée") -> str:
    key = _fold_region_key(value)
    if not key:
        return fallback or "Non renseignée"
    if key in REGION_ALIASES:
        return REGION_ALIASES[key]
    return _normalise_text(value)


def region_threshold(region: str) -> float | None:
    return REGION_NUR_THRESHOLDS.get(region)


def _with_threshold(record: dict) -> dict:
    threshold = region_threshold(record.get("region", ""))
    nur = float(record["nur"])
    record["threshold"] = threshold
    record["above_threshold"] = threshold is not None and nur > threshold
    record["threshold_gap"] = (nur - threshold) if threshold is not None else None
    return record


def _region_sort_key(name: str) -> tuple:
    if name in REGION_ORDER:
        return (0, REGION_ORDER.index(name), name)
    return (1, 99, name)


def top_worst_sites(
    sites: list[dict],
    *,
    region: str | None = None,
    limit: int = 15,
    above_threshold_only: bool = False,
) -> list[dict]:
    selected = [
        row
        for row in sites
        if (region is None or row.get("region") == region)
        and (not above_threshold_only or row.get("above_threshold"))
    ]
    selected.sort(key=lambda row: row["nur"], reverse=True)
    return selected[:limit]


def worst_sites_by_region(sites: list[dict], *, limit: int = 10) -> list[dict]:
    """Les `limit` sites au NUR le plus élevé, regroupés par région."""
    present = {row.get("region") for row in sites}
    names = [name for name in REGION_ORDER if name in present]
    extras = sorted(name for name in present if name not in REGION_ORDER and name)
    ranked: list[dict] = []
    for region in names + extras:
        for rank, row in enumerate(
            top_worst_sites(sites, region=region, limit=limit), start=1
        ):
            ranked.append({"rank": rank, **row})
    return ranked


def parse_nur_date(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = str(value).strip()
    # Les exports Huawei observés utilisent MM/JJ/AAAA.
    for date_format in ("%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    return None


def _header_key(value: object) -> str:
    return _normalise_text(value).casefold().replace(" ", "_")


def _find_header(headers: list[str], *names: str) -> int | None:
    wanted = {_header_key(name) for name in names}
    for index, header in enumerate(headers):
        if _header_key(header) in wanted:
            return index
    return None


def _looks_like_site_code(value: object) -> bool:
    return bool(re.match(r"^[A-Z]{2}_", _normalise_text(value), flags=re.IGNORECASE))


def _sheet_rows(worksheet):
    # Les trois onglets NUR déclarent à tort une dimension A1.
    if worksheet.max_row == 1 and worksheet.max_column == 1:
        worksheet.reset_dimensions()
    rows = worksheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration:
        return [], iter(())
    return [_normalise_text(value) for value in raw_headers], rows


def _site_catalog(
    source: str | Path | BinaryIO,
) -> tuple[set[str], dict[str, str], dict[str, str], str]:
    workbook = _open_workbook(source)
    try:
        if len(workbook.worksheets) < 2:
            raise ValueError(
                "La base des sites doit contenir au moins deux feuilles. "
                "Les données déjà filtrées doivent se trouver dans la deuxième feuille."
            )
        worksheet = workbook.worksheets[1]
        headers, rows = _sheet_rows(worksheet)
        site_index = _find_header(headers, SITE_CODE_HEADER)
        name_index = _find_header(headers, "SITE_NAME", "Site_Name", "Nom_site", "Nom du site")
        if site_index is None and len(headers) > SITE_CODE_COLUMN_INDEX and _looks_like_site_code(
            headers[SITE_CODE_COLUMN_INDEX]
        ):
            # Deuxième feuille ORDC sans ligne d'en-tête : C = nom, D = code, K = région.
            rows = _prepend_row(headers, rows)
            site_index = SITE_CODE_COLUMN_INDEX
            name_index = SITE_NAME_COLUMN_INDEX
            region_index = COMMERCIAL_REGION_COLUMN_INDEX
        elif site_index is None:
            raise ValueError(
                f"La colonne « Site_Code » est absente de la deuxième feuille "
                f"« {worksheet.title} »."
            )
        elif COMMERCIAL_REGION_HEADER in headers or _find_header(headers, COMMERCIAL_REGION_HEADER) is not None:
            region_index = _find_header(headers, COMMERCIAL_REGION_HEADER)
        elif len(headers) > COMMERCIAL_REGION_COLUMN_INDEX:
            region_index = COMMERCIAL_REGION_COLUMN_INDEX
        else:
            raise ValueError(
                "La colonne « Region_commerciale » (colonne K) est absente "
                f"de la deuxième feuille « {worksheet.title} »."
            )
        if region_index is None:
            raise ValueError(
                "La colonne « Region_commerciale » (colonne K) est absente "
                f"de la deuxième feuille « {worksheet.title} »."
            )
        eligible: set[str] = set()
        site_regions: dict[str, str] = {}
        site_names: dict[str, str] = {}
        for row in rows:
            site_code = (
                normalise_site_code(row[site_index])
                if len(row) > site_index
                else ""
            )
            if not site_code:
                continue
            raw_region = (
                row[region_index] if len(row) > region_index else None
            )
            region = classify_region(raw_region)
            if site_code not in site_regions or site_regions[site_code] == "Non renseignée":
                site_regions[site_code] = region
            if name_index is not None and len(row) > name_index:
                site_name = _normalise_text(row[name_index])
                if site_name and not site_names.get(site_code):
                    site_names[site_code] = site_name
            eligible.add(site_code)

        return eligible, site_regions, site_names, worksheet.title
    finally:
        workbook.close()


def _prepend_row(first: list[str], rows):
    yield tuple(first)
    yield from rows


def _float_value(value: object) -> float | None:
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number != number:  # NaN
        return None
    return number


def _read_nur(source: str | Path | BinaryIO) -> tuple[list[dict], list[dict]]:
    workbook = _open_workbook(source)
    records: list[dict] = []
    rejected: list[dict] = []
    try:
        missing_sheets = set(NUR_SHEETS).difference(workbook.sheetnames)
        if missing_sheets:
            raise ValueError(
                "Onglets absents du fichier NUR : " + ", ".join(sorted(missing_sheets))
            )

        for technology, config in NUR_SHEETS.items():
            headers, rows = _sheet_rows(workbook[technology])
            required_headers = {
                config["date"],
                config["site"],
                config["cell"],
                config["unavailability"],
            }
            missing = required_headers.difference(headers)
            if missing:
                raise ValueError(
                    f"Colonnes absentes de l'onglet {technology} : "
                    + ", ".join(sorted(missing))
                )
            indexes = {
                name: headers.index(column)
                for name, column in config.items()
                if name != "suffix"
            }

            for source_row, row in enumerate(rows, start=2):
                def value(name: str):
                    index = indexes[name]
                    return row[index] if len(row) > index else None

                parsed_date = parse_nur_date(value("date"))
                site_code = normalise_site_code(value("site"), config["suffix"])
                cell_name = _normalise_text(value("cell"))
                unavailability = _float_value(value("unavailability"))

                reasons = []
                if parsed_date is None:
                    reasons.append("date invalide")
                if not site_code:
                    reasons.append("site vide")
                if not cell_name:
                    reasons.append("cellule vide")
                if unavailability is None:
                    reasons.append("indisponibilité invalide")
                elif unavailability < 0:
                    reasons.append("indisponibilité négative")

                record = {
                    "date": parsed_date,
                    "technology": technology,
                    "site_code": site_code,
                    "cell_name": cell_name,
                    "unavailability_seconds": unavailability,
                    "source_row": source_row,
                }
                if reasons:
                    rejected.append(
                        {
                            **record,
                            "reason": ", ".join(reasons),
                            "category": "Ligne invalide",
                        }
                    )
                else:
                    records.append(record)
    finally:
        workbook.close()
    return records, rejected


def import_workbooks(
    sites_source: str | Path | BinaryIO,
    nur_source: str | Path | BinaryIO,
) -> ImportResult:
    eligible_sites, site_regions, site_names, database_sheet_name = _site_catalog(sites_source)
    nur_data, rejected = _read_nur(nur_source)
    total_nur_rows = len(nur_data) + len(rejected)
    accepted = []

    for record in nur_data:
        site_code = record["site_code"]
        if site_code not in eligible_sites:
            rejected.append(
                {
                    **record,
                    "reason": "Site absent de la base admissible",
                    "category": "Site non trouvé",
                }
            )
        else:
            catalog_region = site_regions.get(site_code, "Non renseignée")
            accepted.append(
                {
                    **record,
                    "catalog_region": catalog_region,
                    "region": catalog_region,
                    "site_name": site_names.get(site_code, ""),
                }
            )

    return ImportResult(
        data=accepted,
        rejected=rejected,
        eligible_site_count=len(eligible_sites),
        excluded_vsat_site_count=0,
        total_nur_rows=total_nur_rows,
        database_sheet_name=database_sheet_name,
    )


def _nur(unavailability: float, cell_count: int, days: int) -> float:
    return unavailability / (cell_count * SECONDS_PER_DAY * days) * NUR_SCALE


def _add_durations(record: dict) -> dict:
    seconds = float(record["unavailability_seconds"])
    record["unavailability_minutes"] = seconds / 60
    record["unavailability_hours"] = seconds / 3_600
    return record


def calculate_nur(
    data: Iterable[dict],
    start_date: date | datetime,
    period_days: int,
) -> CalculationResult:
    if period_days not in (1, 7, 30):
        raise ValueError("La période doit être de 1, 7 ou 30 jours.")
    start = start_date.date() if isinstance(start_date, datetime) else start_date
    end_exclusive = start + timedelta(days=period_days)
    selected = [row for row in data if start <= row["date"] < end_exclusive]

    site_names: dict[str, str] = {}
    for row in selected:
        site_name = _normalise_text(row.get("site_name"))
        if site_name and not site_names.get(row["site_code"]):
            site_names[row["site_code"]] = site_name

    daily_cells: dict[tuple, float] = {}
    for row in selected:
        region = classify_region(row.get("region", "Non renseignée"))
        key = (
            row["date"],
            row["technology"],
            row["site_code"],
            row["cell_name"],
            region,
        )
        daily_cells[key] = daily_cells.get(key, 0.0) + row[
            "unavailability_seconds"
        ]

    period_cells: dict[tuple, float] = {}
    for (_, technology, site_code, cell_name, region), seconds in daily_cells.items():
        key = (technology, site_code, cell_name, region)
        period_cells[key] = period_cells.get(key, 0.0) + seconds

    cells = [
        _with_threshold(
            _add_durations({
                "technology": technology,
                "site_code": site_code,
                "site_name": site_names.get(site_code, ""),
                "cell_name": cell_name,
                "region": region,
                "unavailability_seconds": seconds,
                "cell_count": 1,
                "period_days": period_days,
                "nur": _nur(seconds, 1, period_days),
            })
        )
        for (technology, site_code, cell_name, region), seconds in period_cells.items()
    ]
    cells.sort(key=lambda row: (row["site_code"], row["technology"], row["cell_name"]))

    technology_groups: dict[tuple, dict] = {}
    site_groups: dict[str, dict] = {}
    for row in cells:
        tech_key = (row["site_code"], row["technology"])
        tech_group = technology_groups.setdefault(
            tech_key,
            {"seconds": 0.0, "cells": set(), "region": row["region"]},
        )
        tech_group["seconds"] += row["unavailability_seconds"]
        tech_group["cells"].add(row["cell_name"])

        site_group = site_groups.setdefault(
            row["site_code"],
            {
                "seconds": 0.0,
                "cells": set(),
                "technologies": set(),
                "region": row["region"],
            },
        )
        site_group["seconds"] += row["unavailability_seconds"]
        site_group["cells"].add((row["technology"], row["cell_name"]))
        site_group["technologies"].add(row["technology"])

    site_technology = []
    for (site_code, technology), group in sorted(technology_groups.items()):
        cell_count = len(group["cells"])
        site_technology.append(
            _with_threshold(
                _add_durations({
                    "site_code": site_code,
                    "site_name": site_names.get(site_code, ""),
                    "technology": technology,
                    "region": group["region"],
                    "unavailability_seconds": group["seconds"],
                    "cell_count": cell_count,
                    "period_days": period_days,
                    "nur": _nur(group["seconds"], cell_count, period_days),
                })
            )
        )

    sites = []
    for site_code, group in sorted(site_groups.items()):
        cell_count = len(group["cells"])
        sites.append(
            _with_threshold(
                _add_durations({
                    "site_code": site_code,
                    "site_name": site_names.get(site_code, ""),
                    "region": group["region"],
                    "unavailability_seconds": group["seconds"],
                    "cell_count": cell_count,
                    "technology_count": len(group["technologies"]),
                    "period_days": period_days,
                    "nur": _nur(group["seconds"], cell_count, period_days),
                })
            )
        )

    technology_totals: dict[str, dict] = {}
    region_totals: dict[str, dict] = {}
    for row in cells:
        technology_group = technology_totals.setdefault(
            row["technology"], {"seconds": 0.0, "cells": set(), "sites": set()}
        )
        technology_group["seconds"] += row["unavailability_seconds"]
        technology_group["cells"].add((row["site_code"], row["cell_name"]))
        technology_group["sites"].add(row["site_code"])

        region_group = region_totals.setdefault(
            row["region"], {"seconds": 0.0, "cells": set(), "sites": set()}
        )
        region_group["seconds"] += row["unavailability_seconds"]
        region_group["cells"].add(
            (row["site_code"], row["technology"], row["cell_name"])
        )
        region_group["sites"].add(row["site_code"])

    technologies = []
    for technology, group in sorted(technology_totals.items()):
        cell_count = len(group["cells"])
        technologies.append(
            _add_durations({
                "technology": technology,
                "site_count": len(group["sites"]),
                "cell_count": cell_count,
                "unavailability_seconds": group["seconds"],
                "period_days": period_days,
                "nur": _nur(group["seconds"], cell_count, period_days),
            })
        )

    regions = []
    for region, group in sorted(region_totals.items(), key=lambda item: _region_sort_key(item[0])):
        cell_count = len(group["cells"])
        regions.append(
            _with_threshold(
                _add_durations({
                    "region": region,
                    "site_count": len(group["sites"]),
                    "cell_count": cell_count,
                    "unavailability_seconds": group["seconds"],
                    "period_days": period_days,
                    "nur": _nur(group["seconds"], cell_count, period_days),
                })
            )
        )

    daily_totals: dict[date, dict] = {}
    for (
        day,
        technology,
        site_code,
        cell_name,
        _region,
    ), seconds in daily_cells.items():
        group = daily_totals.setdefault(
            day, {"seconds": 0.0, "cells": set(), "sites": set()}
        )
        group["seconds"] += seconds
        group["cells"].add((site_code, technology, cell_name))
        group["sites"].add(site_code)

    daily = []
    for day, group in sorted(daily_totals.items()):
        cell_count = len(group["cells"])
        daily.append(
            _add_durations({
                "date": day,
                "site_count": len(group["sites"]),
                "cell_count": cell_count,
                "unavailability_seconds": group["seconds"],
                "period_days": 1,
                "nur": _nur(group["seconds"], cell_count, 1),
            })
        )

    if cells:
        network_cells = {
            (row["site_code"], row["technology"], row["cell_name"]) for row in cells
        }
        total_seconds = sum(row["unavailability_seconds"] for row in cells)
        network = [
            _add_durations({
                "site_count": len({row["site_code"] for row in cells}),
                "technology_count": len({row["technology"] for row in cells}),
                "cell_count": len(network_cells),
                "unavailability_seconds": total_seconds,
                "period_days": period_days,
                "nur": _nur(total_seconds, len(network_cells), period_days),
            })
        ]
    else:
        network = []

    worst_sites = top_worst_sites(sites, limit=50)
    sites_above_threshold = top_worst_sites(
        sites, limit=len(sites), above_threshold_only=True
    )

    return CalculationResult(
        cells=cells,
        site_technology=site_technology,
        sites=sites,
        technologies=technologies,
        regions=regions,
        daily=daily,
        network=network,
        worst_sites=worst_sites,
        sites_above_threshold=sites_above_threshold,
        start_date=start,
        end_date=end_exclusive - timedelta(days=1),
        period_days=period_days,
        selected_rows=len(selected),
    )


EXPORT_LABELS = {
    "date": "Date",
    "technology": "Technologie",
    "region": "Région",
    "source": "Source",
    "site_code": "Code site",
    "site_name": "Nom du site",
    "cell_name": "Cellule",
    "unavailability_seconds": "Indisponibilité (s)",
    "unavailability_minutes": "Indisponibilité (min)",
    "unavailability_hours": "Indisponibilité (h)",
    "source_row": "Ligne source",
    "cell_count": "Nombre de cellules",
    "technology_count": "Nombre de technologies",
    "site_count": "Nombre de sites",
    "period_days": "Nombre de jours",
    "rank": "Rang",
    "nur": "NUR",
    "threshold": "Seuil NUR",
    "above_threshold": "Hors seuil",
    "threshold_gap": "Écart au seuil",
    "reason": "Motif",
    "category": "Catégorie",
}


def _write_records(worksheet, records: list[dict], columns: list[str] | None = None):
    if columns is None:
        columns = list(records[0]) if records else []
    if not columns:
        worksheet.append(["Aucune donnée"])
        return
    worksheet.append([EXPORT_LABELS.get(column, column) for column in columns])
    for record in records:
        worksheet.append([record.get(column) for column in columns])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="F16E00")
    for index, column_cells in enumerate(worksheet.columns, start=1):
        max_length = max(
            (len(str(cell.value)) for cell in column_cells if cell.value is not None),
            default=0,
        )
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            max(max_length + 2, 10), 45
        )


def build_excel_report(calculation: CalculationResult, rejected: list[dict]) -> bytes:
    workbook = Workbook()
    parameters = workbook.active
    parameters.title = "Parametres"
    parameter_rows = [
        {"parameter": "Début", "value": calculation.start_date.isoformat()},
        {"parameter": "Fin", "value": calculation.end_date.isoformat()},
        {"parameter": "Nombre de jours", "value": calculation.period_days},
        {"parameter": "Lignes retenues", "value": calculation.selected_rows},
        {
            "parameter": "Formule",
            "value": "indisponibilité / (cellules × 86400 × jours) × 100000",
        },
    ]
    _write_records(parameters, parameter_rows, ["parameter", "value"])

    threshold_rows = [
        {
            "region": region,
            "source": "Region_commerciale",
            "threshold": threshold,
        }
        for region, threshold in REGION_NUR_THRESHOLDS.items()
    ]
    sheets = [
        ("Cellules", calculation.cells),
        ("Sites_Technologies", calculation.site_technology),
        ("Sites", calculation.sites),
        ("Top_mauvais_sites", calculation.worst_sites[:30]),
        ("Top_10_par_region", worst_sites_by_region(calculation.sites, limit=10)),
        ("Sites_hors_seuil", calculation.sites_above_threshold),
        ("Technologies", calculation.technologies),
        ("Regions", calculation.regions),
        ("Seuils", threshold_rows),
        ("Evolution_journaliere", calculation.daily),
        ("Reseau", calculation.network),
        ("Exclusions", rejected),
    ]
    for name, records in sheets:
        worksheet = workbook.create_sheet(name)
        default_columns = DATA_COLUMNS + ["reason", "category"] if name == "Exclusions" else None
        _write_records(worksheet, records, default_columns)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
