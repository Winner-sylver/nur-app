# -*- coding: utf-8 -*-
"""Présentation NUR — chrome Alarmes, graphes en barres dessinées."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from nur_calculator import (
    REGION_NUR_THRESHOLDS,
    REGION_ORDER,
    CalculationResult,
    ImportResult,
    top_worst_sites,
)

NAVY = RGBColor(0x0B, 0x1F, 0x33)
NAVY2 = RGBColor(0x14, 0x32, 0x4A)
ORANGE = RGBColor(0xFF, 0x79, 0x00)
BG = RGBColor(0xF6, 0xF3, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x0F, 0x7B, 0x4A)
AMBER = RGBColor(0xC4, 0x6B, 0x00)
RED = RGBColor(0xC4, 0x1E, 0x3A)
GRAY = RGBColor(0x5C, 0x67, 0x73)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GREEN_BG = RGBColor(0xE6, 0xF4, 0xEA)
AMBER_BG = RGBColor(0xFF, 0xF3, 0xD6)
RED_BG = RGBColor(0xF8, 0xD4, 0xD0)
ALT = RGBColor(0xFB, 0xF7, 0xF2)
TRACK = RGBColor(0xE6, 0xE1, 0xD8)
MUTED = RGBColor(0xC8, 0xD0, 0xD8)

FONT = "Calibri"
W, H = Inches(13.333), Inches(7.5)
TOTAL = 8
TECH_COLORS = {"2G": GRAY, "3G": NAVY2, "4G": ORANGE}


def _fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _rect(slide, x, y, w, h, color, *, rounded=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, x, y, w, h)
    _fill(shape, color)
    if rounded:
        try:
            shape.adjustments[0] = 0.06
        except Exception:
            pass
    return shape


def _run(run, *, size=14, bold=False, color=BLACK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT


def _tb(slide, x, y, w, h, text, *, size=14, bold=False, color=BLACK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _run(run, size=size, bold=bold, color=color)
    return box


def _content_slide(prs, kicker, title, page, *, foot=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, Inches(0), Inches(0), W, H, BG)
    _rect(slide, Inches(0), Inches(0), W, Inches(0.08), ORANGE)
    _rect(slide, Inches(0), Inches(0.08), W, Inches(0.86), NAVY)
    _tb(slide, Inches(0.48), Inches(0.14), Inches(12.3), Inches(0.22), kicker, size=10, bold=True, color=ORANGE)
    _tb(slide, Inches(0.48), Inches(0.36), Inches(12.3), Inches(0.48), title, size=22, bold=True, color=WHITE)
    _rect(slide, Inches(0), Inches(7.22), W, Inches(0.28), NAVY)
    _tb(slide, Inches(0.48), Inches(7.24), Inches(9.2), Inches(0.24), foot, size=10, color=WHITE)
    _tb(
        slide, Inches(10.5), Inches(7.24), Inches(2.35), Inches(0.24),
        f"{page}  /  {TOTAL}", size=10, color=WHITE, align=PP_ALIGN.RIGHT,
    )
    return slide


def _kpi(slide, x, y, w, h, label, value, hint="", *, accent=ORANGE, vcolor=BLACK):
    _rect(slide, x, y, w, h, WHITE, rounded=True)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.08), h)
    _fill(bar, accent)
    _tb(slide, x + Inches(0.2), y + Inches(0.08), w - Inches(0.28), Inches(0.22), label.upper(), size=10, bold=True, color=GRAY)
    _tb(slide, x + Inches(0.2), y + Inches(0.3), w - Inches(0.28), Inches(0.42), value, size=22, bold=True, color=vcolor)
    if hint:
        _tb(slide, x + Inches(0.2), y + h - Inches(0.32), w - Inches(0.28), Inches(0.26), hint, size=11, color=GRAY)


def _fill_cell(cell, text, *, size=11, bold=False, color=BLACK, bg=None, align=PP_ALIGN.CENTER):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    _run(run, size=size, bold=bold, color=color)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.05)
    cell.margin_right = Inches(0.04)
    cell.margin_top = Inches(0.02)
    cell.margin_bottom = Inches(0.02)
    if bg is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg
    else:
        cell.fill.background()


def _borders(table, hexcol="E6E1D8"):
    for cell in table.iter_cells():
        tc_pr = cell._tc.get_or_add_tcPr()
        for tag in ("lnL", "lnR", "lnT", "lnB"):
            ln = tc_pr.find(qn(f"a:{tag}"))
            if ln is None:
                ln = etree.SubElement(tc_pr, qn(f"a:{tag}"))
            ln.set("w", "6350")
            for child in list(ln):
                ln.remove(child)
            solid = etree.SubElement(ln, qn("a:solidFill"))
            srgb = etree.SubElement(solid, qn("a:srgbClr"))
            srgb.set("val", hexcol)


def _site_title(row: dict) -> str:
    name = str(row.get("site_name") or "").strip()
    code = str(row.get("site_code") or "").strip()
    return name or code or "—"


def _site_caption(row: dict) -> str:
    name = str(row.get("site_name") or "").strip()
    code = str(row.get("site_code") or "").strip()
    if name and name.casefold() != code.casefold():
        return f"{name}  ·  {code}"
    return code or "—"


def _fmt(value: float) -> str:
    if value >= 1000:
        return f"{value:,.0f}".replace(",", " ")
    if value >= 10:
        return f"{value:,.1f}".replace(",", " ")
    return f"{value:.2f}"


def _status_colors(above: bool):
    return (RED_BG, RED) if above else (GREEN_BG, GREEN)


def _bar_row(slide, x, y, label_w, track_w, label, value, *, maximum, threshold=None):
    above = threshold is not None and value > threshold
    color = RED if above else ORANGE
    _tb(slide, x, y, label_w, Inches(0.28), label, size=11, color=NAVY)
    _rect(slide, x + label_w, y + Inches(0.06), track_w, Inches(0.16), TRACK, rounded=True)
    ratio = 0 if maximum <= 0 else min(value / maximum, 1)
    if ratio > 0:
        fill_w = max(Emu(28000), int(track_w * ratio))
        bar = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            x + label_w,
            y + Inches(0.06),
            fill_w,
            Inches(0.16),
        )
        _fill(bar, color)
    if threshold is not None and maximum > 0:
        marker = x + label_w + int(track_w * min(threshold / maximum, 1))
        tick = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, marker, y + Inches(0.02), Inches(0.018), Inches(0.24)
        )
        _fill(tick, NAVY)
    _tb(
        slide,
        x + label_w + track_w + Inches(0.08),
        y,
        Inches(0.95),
        Inches(0.28),
        _fmt(value),
        size=11,
        bold=True,
        color=color,
    )


def _style_line(chart, ymin=None, ymax=None):
    chart.has_title = False
    chart.has_legend = False
    try:
        chart.font.size = Pt(10)
        chart.font.name = FONT
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.tick_labels.font.size = Pt(9)
        chart.category_axis.tick_labels.font.size = Pt(9)
    except Exception:
        pass
    if ymin is None and ymax is None:
        return
    scaling = chart.value_axis._element.find(qn("c:scaling"))
    if scaling is None:
        scaling = etree.SubElement(chart.value_axis._element, qn("c:scaling"))
    if ymin is not None:
        el = scaling.find(qn("c:min"))
        if el is None:
            el = etree.SubElement(scaling, qn("c:min"))
        el.set("val", str(ymin))
    if ymax is not None:
        el = scaling.find(qn("c:max"))
        if el is None:
            el = etree.SubElement(scaling, qn("c:max"))
        el.set("val", str(ymax))


def build_powerpoint_report(
    calculation: CalculationResult, imported: ImportResult
) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    network = calculation.network[0] if calculation.network else {
        "nur": 0, "site_count": 0, "cell_count": 0, "unavailability_hours": 0
    }
    period = (
        f"{calculation.start_date.strftime('%d/%m/%Y')} – "
        f"{calculation.end_date.strftime('%d/%m/%Y')}"
    )
    foot = f"SMC  ·  Observatoire NUR  ·  {period}  ·  {calculation.period_days} j"
    worst = calculation.worst_sites or top_worst_sites(calculation.sites, limit=50)
    hors_seuil = calculation.sites_above_threshold
    worst_site = worst[0] if worst else None
    regions = list(calculation.regions)

    # 1 COVER
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, Inches(0), Inches(0), W, H, NAVY)
    _rect(slide, Inches(0), Inches(0), Inches(0.18), H, ORANGE)
    _rect(slide, Inches(0), Inches(6.85), W, Inches(0.65), ORANGE)
    _tb(
        slide, Inches(0.85), Inches(1.1), Inches(11), Inches(0.3),
        "SERVICE MANAGEMENT CENTER  ·  OBSERVATOIRE NUR",
        size=12, bold=True, color=ORANGE,
    )
    _tb(slide, Inches(0.85), Inches(1.55), Inches(11.5), Inches(1.15), "NUR réseau", size=40, bold=True, color=WHITE)
    _tb(
        slide, Inches(0.85), Inches(2.8), Inches(11.2), Inches(0.5),
        f"{period}  —  2G · 3G · 4G  ·  seuils par région  ·  top sites",
        size=14, color=MUTED,
    )
    metas = [
        ("PÉRIODE", f"{calculation.period_days} jour(s)"),
        ("SITES", f"{network['site_count']}"),
        ("HORS SEUIL", f"{len(hors_seuil)}"),
        ("NUR RÉSEAU", _fmt(network["nur"])),
    ]
    for i, (a, b) in enumerate(metas):
        x = Inches(0.85 + i * 3.05)
        _rect(slide, x, Inches(4.15), Inches(2.88), Inches(1.2), NAVY2, rounded=True)
        _tb(slide, x + Inches(0.16), Inches(4.28), Inches(2.56), Inches(0.28), a, size=11, bold=True, color=ORANGE)
        _tb(slide, x + Inches(0.16), Inches(4.6), Inches(2.56), Inches(0.55), str(b), size=18, bold=True, color=WHITE)
    _tb(
        slide, Inches(0.85), Inches(6.98), Inches(11.5), Inches(0.38),
        f"Document de pilotage  ·  {imported.database_sheet_name}  ·  {datetime.now().strftime('%d/%m/%Y')}",
        size=12, bold=True, color=NAVY,
    )

    # 2 EXEC
    slide = _content_slide(prs, "01 / VUE EXÉCUTIVE", "Ce qu’il faut retenir", 2, foot=foot)
    _kpi(slide, Inches(0.4), Inches(1.12), Inches(3.05), Inches(1.2), "NUR réseau", _fmt(network["nur"]), "Tous les sites", accent=ORANGE)
    _kpi(
        slide, Inches(3.58), Inches(1.12), Inches(3.05), Inches(1.2),
        "Hors seuil", str(len(hors_seuil)), "Sites au-dessus du seuil régional",
        accent=RED if hors_seuil else GREEN,
        vcolor=RED if hors_seuil else GREEN,
    )
    _kpi(
        slide, Inches(6.76), Inches(1.12), Inches(3.05), Inches(1.2),
        "Sites", str(network["site_count"]),
        f"{network['cell_count']} cellules",
        accent=NAVY,
        vcolor=NAVY,
    )
    top_label = worst_site["site_code"] if worst_site else "—"
    top_hint = (
        f"{worst_site['region']}  ·  NUR {_fmt(worst_site['nur'])}"
        if worst_site else "Aucun site"
    )
    _kpi(
        slide, Inches(9.94), Inches(1.12), Inches(2.9), Inches(1.2),
        "Pire site", top_label, top_hint, accent=AMBER, vcolor=AMBER,
    )

    for i, tech in enumerate(sorted(calculation.technologies, key=lambda row: row["technology"])):
        x = Inches(0.4 + i * 4.25)
        color = TECH_COLORS.get(tech["technology"], ORANGE)
        _rect(slide, x, Inches(2.52), Inches(4.05), Inches(1.55), WHITE, rounded=True)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(2.52), Inches(0.08), Inches(1.55))
        _fill(bar, color)
        _tb(slide, x + Inches(0.24), Inches(2.68), Inches(3.6), Inches(0.28), tech["technology"], size=14, bold=True, color=color)
        _tb(slide, x + Inches(0.24), Inches(3.02), Inches(3.6), Inches(0.45), _fmt(tech["nur"]), size=22, bold=True, color=NAVY)
        _tb(
            slide, x + Inches(0.24), Inches(3.5), Inches(3.6), Inches(0.35),
            f"{tech['site_count']} sites  ·  {tech['cell_count']} cellules",
            size=12, color=GRAY,
        )

    _rect(slide, Inches(0.4), Inches(4.3), Inches(12.5), Inches(2.7), WHITE, rounded=True)
    _tb(slide, Inches(0.62), Inches(4.48), Inches(12), Inches(0.3), "Lecture", size=14, bold=True, color=ORANGE)
    worst_txt = (
        f"Pire site : {worst_site['site_code']} ({worst_site['region']}), NUR {_fmt(worst_site['nur'])}"
        + (
            f"  ·  seuil {int(worst_site['threshold'])}."
            if worst_site and worst_site.get("threshold")
            else "."
        )
        if worst_site else "Aucun site classé."
    )
    _tb(
        slide, Inches(0.62), Inches(4.9), Inches(12), Inches(1.85),
        f"{len(hors_seuil)} site(s) dépassent le seuil régional sur {network['site_count']}.\n"
        f"{worst_txt}\n"
        "Préfixes : KN Kinshasa  ·  KT Katanga  ·  KO/KE Kasaï  ·  KC Kongo Central  ·  KV Kivu  ·  PO/EQ Grand Nord.",
        size=14, color=NAVY,
    )

    # 3 REGIONS vs SEUILS
    slide = _content_slide(
        prs, "02 / RÉGIONS", "NUR régional confronté au seuil", 3, foot=foot
    )
    max_region = max(
        [row["nur"] for row in regions] + list(REGION_NUR_THRESHOLDS.values()),
        default=1,
    )
    _tb(
        slide, Inches(0.48), Inches(1.08), Inches(12.3), Inches(0.28),
        "Barre orange = sous le seuil   ·   rouge = hors seuil   ·   trait navy = seuil de la région",
        size=12, color=GRAY,
    )
    for i, row in enumerate(regions[:8]):
        y = Inches(1.42 + i * 0.62)
        label = f"{row['region']}  ({int(row['threshold']) if row.get('threshold') else '—'})"
        _bar_row(
            slide, Inches(0.48), y, Inches(3.35), Inches(7.7),
            label, row["nur"], maximum=max_region * 1.05, threshold=row.get("threshold"),
        )

    # 4 TOP SITES
    slide = _content_slide(
        prs, "03 / TOP MAUVAIS SITES", "Les 12 NUR les plus élevés de la période", 4, foot=foot
    )
    top12 = worst[:12]
    max_site = max((row["nur"] for row in top12), default=1)
    for i, row in enumerate(top12):
        y = Inches(1.1 + i * 0.48)
        label = _site_caption(row) + f"  ·  {row['region']}"
        _bar_row(
            slide, Inches(0.4), y, Inches(4.15), Inches(6.9),
            label, row["nur"], maximum=max_site * 1.08, threshold=row.get("threshold"),
        )

    # 5 TOP 10 PAR RÉGION
    slide = _content_slide(
        prs, "04 / PAR RÉGION", "Les 10 pires sites de chaque région", 5, foot=foot
    )
    shown_regions = [name for name in REGION_ORDER if any(row["region"] == name for row in calculation.sites)]
    col_w = Inches(2.08)
    for i, region in enumerate(shown_regions[:6]):
        x = Inches(0.32 + i * 2.16)
        threshold = REGION_NUR_THRESHOLDS[region]
        _rect(slide, x, Inches(1.12), col_w, Inches(5.95), WHITE, rounded=True)
        _tb(slide, x + Inches(0.08), Inches(1.18), col_w - Inches(0.16), Inches(0.28), region, size=11, bold=True, color=NAVY)
        _tb(
            slide, x + Inches(0.08), Inches(1.42), col_w - Inches(0.16), Inches(0.2),
            f"seuil {threshold}", size=9, bold=True, color=ORANGE,
        )
        local = top_worst_sites(calculation.sites, region=region, limit=10)
        if not local:
            _tb(slide, x + Inches(0.08), Inches(2.0), col_w - Inches(0.16), Inches(0.3), "Aucun site", size=10, color=GRAY)
            continue
        for j, site in enumerate(local):
            tone = RED if site.get("above_threshold") else NAVY
            yy = Inches(1.68 + j * 0.52)
            _tb(
                slide, x + Inches(0.08), yy, Inches(1.28), Inches(0.18),
                f"{j + 1}. {_site_title(site)}", size=8, bold=True, color=tone,
            )
            if str(site.get("site_name") or "").strip():
                _tb(
                    slide, x + Inches(0.18), yy + Inches(0.16), Inches(1.15), Inches(0.16),
                    site["site_code"], size=7, color=GRAY,
                )
            _tb(
                slide, x + Inches(1.22), yy, Inches(0.78), Inches(0.18),
                _fmt(site["nur"]), size=8, bold=True, color=tone, align=PP_ALIGN.RIGHT,
            )

    # 6 TABLE HORS SEUIL
    slide = _content_slide(
        prs, "05 / HORS SEUIL", "Sites au-dessus du seuil régional", 6, foot=foot
    )
    rows = hors_seuil[:14]
    if not rows:
        _rect(slide, Inches(0.48), Inches(1.3), Inches(12.37), Inches(5.5), WHITE, rounded=True)
        _tb(
            slide, Inches(0.8), Inches(3.2), Inches(11.7), Inches(0.8),
            "Aucun site au-dessus du seuil sur cette période.",
            size=18, bold=True, color=GREEN, align=PP_ALIGN.CENTER,
        )
    else:
        table = slide.shapes.add_table(
            len(rows) + 1, 6, Inches(0.4), Inches(1.1), Inches(12.5), Inches(5.9)
        ).table
        widths = [2.3, 1.9, 1.6, 2.0, 2.2, 2.3]
        for j, width in enumerate(widths):
            table.columns[j].width = Inches(width)
        headers = ["Site", "Région", "NUR", "Seuil", "Écart", "Statut"]
        for j, header in enumerate(headers):
            _fill_cell(table.cell(0, j), header, size=11, bold=True, color=WHITE, bg=NAVY)
        for i, row in enumerate(rows, start=1):
            gap = row.get("threshold_gap") or 0
            _fill_cell(table.cell(i, 0), row["site_code"], size=11, bold=True, bg=WHITE if i % 2 else ALT, align=PP_ALIGN.LEFT)
            _fill_cell(table.cell(i, 1), row["region"], size=11, bg=WHITE if i % 2 else ALT, align=PP_ALIGN.LEFT)
            _fill_cell(table.cell(i, 2), _fmt(row["nur"]), size=11, bold=True, bg=WHITE if i % 2 else ALT)
            _fill_cell(table.cell(i, 3), str(int(row["threshold"])) if row.get("threshold") else "—", size=11, bg=WHITE if i % 2 else ALT)
            _fill_cell(table.cell(i, 4), f"+{_fmt(gap)}", size=11, bold=True, color=RED, bg=RED_BG)
            _fill_cell(table.cell(i, 5), "Hors seuil", size=11, bold=True, color=RED, bg=RED_BG)
        _borders(table)

    # 7 ÉVOLUTION
    slide = _content_slide(
        prs, "06 / TENDANCE", "Évolution journalière du NUR réseau", 7, foot=foot
    )
    daily = calculation.daily
    if len(daily) >= 2:
        data = CategoryChartData()
        data.categories = [row["date"].strftime("%d/%m") for row in daily]
        data.add_series("NUR", [round(row["nur"], 2) for row in daily])
        frame = slide.shapes.add_chart(
            XL_CHART_TYPE.LINE, Inches(0.35), Inches(1.15), Inches(8.4), Inches(5.75), data
        )
        chart = frame.chart
        _style_line(chart, ymin=0)
        series = chart.series[0]
        series.format.line.color.rgb = ORANGE
        try:
            series.format.line.width = Pt(2.5)
        except Exception:
            pass
    elif daily:
        _rect(slide, Inches(0.4), Inches(1.2), Inches(8.3), Inches(5.7), WHITE, rounded=True)
        _tb(slide, Inches(0.7), Inches(3.2), Inches(7.7), Inches(0.5), _fmt(daily[0]["nur"]), size=32, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)
        _tb(slide, Inches(0.7), Inches(3.8), Inches(7.7), Inches(0.4), daily[0]["date"].strftime("%d/%m/%Y"), size=14, color=GRAY, align=PP_ALIGN.CENTER)
    _rect(slide, Inches(8.9), Inches(1.2), Inches(3.95), Inches(5.7), WHITE, rounded=True)
    _tb(slide, Inches(9.1), Inches(1.4), Inches(3.55), Inches(0.35), "Seuils", size=16, bold=True, color=NAVY)
    for i, (region, threshold) in enumerate(REGION_NUR_THRESHOLDS.items()):
        yy = Inches(1.9 + i * 0.72)
        _tb(slide, Inches(9.1), yy, Inches(2.2), Inches(0.28), region, size=12, color=NAVY)
        _tb(slide, Inches(11.15), yy, Inches(1.4), Inches(0.28), str(threshold), size=14, bold=True, color=ORANGE, align=PP_ALIGN.RIGHT)

    # 8 MÉTHODE
    slide = _content_slide(
        prs, "07 / MÉTHODE", "Classification, formule et seuils", 8, foot=foot
    )
    _rect(slide, Inches(0.4), Inches(1.14), Inches(6.15), Inches(5.8), WHITE, rounded=True)
    _tb(slide, Inches(0.62), Inches(1.32), Inches(5.7), Inches(0.3), "Région commerciale → seuil", size=16, bold=True, color=NAVY)
    mapping = [
        ("Kinshasa", "800"),
        ("Grand Katanga", "600"),
        ("Grand Kasai", "500"),
        ("Kongo Central", "300"),
        ("Grand Kivu", "1 000"),
        ("Grand Nord", "1 000"),
    ]
    for i, (region, threshold) in enumerate(mapping):
        yy = Inches(1.8 + i * 0.78)
        _tb(slide, Inches(0.7), yy, Inches(3.4), Inches(0.4), region, size=14, bold=True, color=NAVY)
        _tb(slide, Inches(4.3), yy, Inches(1.9), Inches(0.4), threshold, size=14, bold=True, color=ORANGE)
    _rect(slide, Inches(6.8), Inches(1.14), Inches(6.05), Inches(5.8), WHITE, rounded=True)
    _tb(slide, Inches(7.02), Inches(1.32), Inches(5.6), Inches(0.3), "Formule", size=16, bold=True, color=NAVY)
    _tb(
        slide, Inches(7.02), Inches(1.8), Inches(5.6), Inches(2.4),
        "NUR = indisponibilité\n/ (cellules × 86 400 × jours)\n× 100 000",
        size=16, color=NAVY,
    )
    _tb(
        slide, Inches(7.02), Inches(4.4), Inches(5.6), Inches(2.2),
        f"Base : {imported.database_sheet_name}\n"
        f"Lignes retenues : {calculation.selected_rows}\n"
        f"Lignes exclues : {len(imported.rejected)}\n"
        "Un site est « mauvais » s’il dépasse le seuil de sa région.",
        size=14, color=GRAY,
    )

    output = BytesIO()
    prs.save(output)
    return output.getvalue()
