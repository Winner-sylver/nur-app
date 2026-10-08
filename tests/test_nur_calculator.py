from datetime import date
from io import BytesIO
import unittest

from openpyxl import Workbook, load_workbook
from pptx import Presentation

from nur_calculator import (
    REGION_NUR_THRESHOLDS,
    build_excel_report,
    calculate_nur,
    classify_region,
    import_workbooks,
    normalise_site_code,
    parse_nur_date,
    top_worst_sites,
    worst_sites_by_region,
)
from ppt_report import build_powerpoint_report


def workbook_bytes(workbook: Workbook) -> BytesIO:
    content = BytesIO()
    workbook.save(content)
    content.seek(0)
    return content


def make_sites_workbook() -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Instructions"
    sheet.append(["Cette feuille n'est pas utilisée"])
    sheet = workbook.create_sheet("Sites filtres")
    sheet.append(["Site_Code", "Type_transmission", "Region_commerciale", "SITE_NAME"])
    sheet.append(["KN_01_001_C0", "MW", "KINSHASA", "Gombe"])
    sheet.append(["KN_01_002_C0", "MW-Vsat", "Grand Kivu", "Bukavu Centre"])
    return workbook_bytes(workbook)


def make_nur_workbook() -> BytesIO:
    workbook = Workbook()
    sheet_2g = workbook.active
    sheet_2g.title = "2G"
    sheet_2g.append(
        [
            "Date",
            "Cell Name",
            "Site Name",
            "R373:Cell Out-of-Service Duration(s)",
        ]
    )
    sheet_2g.append(["09/03/2026", "CELL-A", "KN_01_001_C0_GSM", 100])
    sheet_2g.append(["09/03/2026", "CELL-B", "KN_01_002_C0_GSM", 200])
    sheet_2g.append(["09/03/2026", "CELL-C", "INCONNU_GSM", 300])

    sheet_3g = workbook.create_sheet("3G")
    sheet_3g.append(
        [
            "Date",
            "Cell Name",
            "NODEBNAME",
            "VS.Cell.UnavailTime.Sys(s)",
        ]
    )
    sheet_4g = workbook.create_sheet("4G")
    sheet_4g.append(
        [
            "Date",
            "Cell Name",
            "eNodeB Name",
            "L.Cell.Unavail.Dur.Sys(s)",
        ]
    )
    return workbook_bytes(workbook)


class NormalisationTests(unittest.TestCase):
    def test_site_suffixes_are_removed(self):
        self.assertEqual(
            normalise_site_code(" kn_01_001_c0_gsm ", r"_GSM$"),
            "KN_01_001_C0",
        )
        self.assertEqual(
            normalise_site_code("KN_01_001_C0_UMTS", r"_UMTS$"),
            "KN_01_001_C0",
        )

    def test_huawei_date_is_month_day_year(self):
        self.assertEqual(parse_nur_date("09/03/2026"), date(2026, 9, 3))

    def test_second_sheet_is_used_without_vsat_filter(self):
        result = import_workbooks(make_sites_workbook(), make_nur_workbook())
        self.assertEqual(len(result.data), 2)
        self.assertEqual(result.data[0]["site_code"], "KN_01_001_C0")
        self.assertEqual(result.data[0]["site_name"], "Gombe")
        self.assertEqual(result.data[0]["region"], "Kinshasa")
        self.assertEqual(result.data[1]["site_code"], "KN_01_002_C0")
        self.assertEqual(result.data[1]["region"], "Grand Kivu")
        categories = {row["category"] for row in result.rejected}
        self.assertEqual(categories, {"Site non trouvé"})
        self.assertEqual(result.database_sheet_name, "Sites filtres")
        self.assertEqual(result.excluded_vsat_site_count, 0)

    def test_headerless_second_sheet_reads_name_and_code(self):
        workbook = Workbook()
        workbook.active.title = "Vide"
        workbook.active.append(["x"])
        sheet = workbook.create_sheet("Filtre")
        row = [""] * 11
        row[2] = "Av.Mpeti"
        row[3] = "KN_01_001_C0"
        row[10] = "KINSHASA"
        sheet.append(row)
        result = import_workbooks(workbook_bytes(workbook), make_nur_workbook())
        self.assertEqual(result.data[0]["site_code"], "KN_01_001_C0")
        self.assertEqual(result.data[0]["site_name"], "Av.Mpeti")
        self.assertEqual(result.data[0]["region"], "Kinshasa")

    def test_database_without_second_sheet_is_rejected(self):
        workbook = Workbook()
        workbook.active.append(["Site_Code"])
        with self.assertRaisesRegex(ValueError, "deux feuilles"):
            import_workbooks(workbook_bytes(workbook), make_nur_workbook())


class CalculationTests(unittest.TestCase):
    def setUp(self):
        self.rows = [
            {
                "date": date(2026, 9, 3),
                "technology": "2G",
                "site_code": "SITE-A",
                "site_name": "Site Alpha",
                "cell_name": "CELL-1",
                "region": "REGION-A",
                "unavailability_seconds": 86_400.0,
            },
            {
                "date": date(2026, 9, 3),
                "technology": "2G",
                "site_code": "SITE-A",
                "cell_name": "CELL-2",
                "region": "REGION-A",
                "unavailability_seconds": 0.0,
            },
            {
                "date": date(2026, 9, 3),
                "technology": "3G",
                "site_code": "SITE-A",
                "cell_name": "CELL-1",
                "region": "REGION-A",
                "unavailability_seconds": 86_400.0,
            },
        ]

    def test_cell_and_technology_daily_nur(self):
        result = calculate_nur(self.rows, date(2026, 9, 3), 1)
        first_cell = next(
            row
            for row in result.cells
            if row["technology"] == "2G" and row["cell_name"] == "CELL-1"
        )
        technology_2g = next(
            row for row in result.site_technology if row["technology"] == "2G"
        )
        self.assertAlmostEqual(first_cell["nur"], 100_000)
        self.assertAlmostEqual(technology_2g["nur"], 50_000)

    def test_site_counts_same_cell_name_in_two_technologies(self):
        result = calculate_nur(self.rows, date(2026, 9, 3), 1)
        site = result.sites[0]
        self.assertEqual(site["cell_count"], 3)
        self.assertEqual(site["technology_count"], 2)
        self.assertAlmostEqual(site["nur"], 200_000 / 3)

    def test_fixed_week_and_month_denominators(self):
        week = calculate_nur(self.rows[:1], date(2026, 9, 3), 7)
        month = calculate_nur(self.rows[:1], date(2026, 9, 3), 30)
        self.assertAlmostEqual(week.network[0]["nur"], 100_000 / 7)
        self.assertAlmostEqual(month.network[0]["nur"], 100_000 / 30)

    def test_durations_and_regional_summary(self):
        result = calculate_nur(self.rows, date(2026, 9, 3), 1)
        self.assertAlmostEqual(result.network[0]["unavailability_hours"], 48)
        self.assertAlmostEqual(result.network[0]["unavailability_minutes"], 2_880)
        self.assertEqual(result.regions[0]["region"], "REGION-A")
        self.assertEqual(result.sites[0]["site_name"], "Site Alpha")
        self.assertEqual(result.regions[0]["site_count"], 1)
        self.assertEqual(len(result.technologies), 2)

    def test_duplicate_daily_rows_are_consolidated(self):
        duplicate = dict(self.rows[0])
        duplicate["unavailability_seconds"] = 86_400.0
        result = calculate_nur(
            [self.rows[0], duplicate], date(2026, 9, 3), 1
        )
        self.assertEqual(len(result.cells), 1)
        self.assertAlmostEqual(result.cells[0]["nur"], 200_000)

    def test_excel_report_contains_all_sheets(self):
        result = calculate_nur(self.rows, date(2026, 9, 3), 1)
        report = build_excel_report(result, [])
        workbook = load_workbook(BytesIO(report), read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            [
                "Parametres",
                "Cellules",
                "Sites_Technologies",
                "Sites",
                "Top_mauvais_sites",
                "Top_10_par_region",
                "Sites_hors_seuil",
                "Technologies",
                "Regions",
                "Seuils",
                "Evolution_journaliere",
                "Reseau",
                "Exclusions",
            ],
        )

    def test_powerpoint_report_is_readable(self):
        imported = import_workbooks(make_sites_workbook(), make_nur_workbook())
        result = calculate_nur(imported.data, date(2026, 9, 3), 1)
        report = build_powerpoint_report(result, imported)
        presentation = Presentation(BytesIO(report))
        self.assertEqual(len(presentation.slides), 8)
        titles = [
            shape.text
            for shape in presentation.slides[1].shapes
            if hasattr(shape, "text") and shape.text
        ]
        self.assertIn("01 / VUE EXÉCUTIVE", titles)

    def test_commercial_region_names_are_normalised(self):
        mapping = {
            "KINSHASA": "Kinshasa",
            "Grand Katanga": "Grand Katanga",
            "GRAND KIVU": "Grand Kivu",
            "Grand Kasai": "Grand Kasai",
            "KASAI BANDUNDU": "Grand Kasai",
            "kongo central": "Kongo Central",
            "Grand Nord": "Grand Nord",
            "": "REGION-A",
        }
        for raw, region in mapping.items():
            self.assertEqual(classify_region(raw, "REGION-A"), region)

    def test_regional_threshold_flags_worst_sites(self):
        rows = [
            {
                "date": date(2026, 9, 3),
                "technology": "2G",
                "site_code": "KN_01_900_C0",
                "cell_name": "CELL-1",
                "region": "KINSHASA",
                "unavailability_seconds": 1_000.0,
            },
            {
                "date": date(2026, 9, 3),
                "technology": "2G",
                "site_code": "KC_01_100_C0",
                "cell_name": "CELL-1",
                "region": "Kongo Central",
                "unavailability_seconds": 100.0,
            },
        ]
        result = calculate_nur(rows, date(2026, 9, 3), 1)
        kinshasa = next(row for row in result.sites if row["site_code"].startswith("KN"))
        kongo = next(row for row in result.sites if row["site_code"].startswith("KC"))
        self.assertEqual(kinshasa["region"], "Kinshasa")
        self.assertEqual(kinshasa["threshold"], 800)
        self.assertGreater(kinshasa["nur"], 800)
        self.assertTrue(kinshasa["above_threshold"])
        self.assertEqual(kongo["region"], "Kongo Central")
        self.assertEqual(kongo["threshold"], 300)
        self.assertLess(kongo["nur"], 300)
        self.assertFalse(kongo["above_threshold"])
        self.assertEqual(result.sites_above_threshold[0]["site_code"], "KN_01_900_C0")
        self.assertEqual(top_worst_sites(result.sites, region="Kinshasa", limit=1)[0]["site_code"], "KN_01_900_C0")
        by_region = worst_sites_by_region(result.sites, limit=10)
        self.assertEqual([row["region"] for row in by_region], ["Kinshasa", "Kongo Central"])
        self.assertEqual([row["rank"] for row in by_region], [1, 1])
        self.assertEqual(REGION_NUR_THRESHOLDS["Grand Kivu"], 1000)


if __name__ == "__main__":
    unittest.main()
